import pandas as pd, numpy as np, re, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter as L
import os; c=pd.read_csv(os.environ.get('RES','results')+'/per_cell_metrics.csv'); c['has_junction']=c.has_junction.astype(int)
OUT=os.environ.get('OUTD','timeline_xlsx'); os.makedirs(OUT,exist_ok=True)
Y=PatternFill('solid',fgColor='FFFF00'); B=Font(name='Calibri',size=11,bold=True); N=Font(name='Calibri',size=11)
def key(n):
    m=re.search(r'_(\d+)$',n); return (n[:m.start()] if m else n, int(m.group(1)) if m else -1)
TIT={'junction_density':'E-cad Junction Density per Cell (µm per 100 µm²)','ridge_snr':'E-cad Ridge SNR per Cell','continuity':'E-cad Junction Continuity per Cell (cells with junctions only)','JC_ratio':'E-cad Junction/Cytoplasm Ratio per Cell (cells with junctions only)','has_junction':'Cell has a detected E-cad junction (1 = yes, 0 = no); Mean = fraction of cells'}
groups=[(tr,t) for tr in ['Control','CalA'] for t in [0,2,5,10,15]]
for metric,title in TIT.items():
    wb=Workbook(); wb.remove(wb.active); info={}
    for tr,t in groups:
        ws=wb.create_sheet(f'{tr} {t} Min'); ws['A1']=f'{title} - {"Calyculin A" if tr=="CalA" else "Control"}, {t} min post wounding'; ws['A1'].font=N
        for col,h in zip('ABD',['No.','Sample ID','Mean']): ws[f'{col}4']=h
        for col in 'ABCD': ws[f'{col}4'].font=B; ws[f'{col}4'].fill=Y
        ws['D4'].alignment=Alignment(horizontal='center')
        d=c[(c.treatment==tr)&(c.time_min==t)]; imgs=sorted(d.image.unique(),key=key); r=5
        for i,im in enumerate(imgs,1):
            v=d[d.image==im][metric].dropna().tolist()
            ws.cell(r,1,i).font=N; ws.cell(r,2,im).font=N
            x=ws.cell(r,4,f'=IFERROR(AVERAGE(E{r}:AAA{r}),"")'); x.number_format='0.00'; x.alignment=Alignment(horizontal='center'); x.font=N
            for j,val in enumerate(v): y=ws.cell(r,5+j,float(val)); y.number_format='0.00'; y.font=N
            r+=1
        info[(tr,t)]=(ws.title,5,r-1)
        ws.column_dimensions['A'].width=5; ws.column_dimensions['B'].width=36; ws.column_dimensions['C'].width=11; ws.column_dimensions['D'].width=12
    ws=wb.create_sheet('Regression')
    for j,(tr,t) in enumerate(groups,1):
        ws.cell(1,j,f'{tr} {t} Min').font=B; ws.column_dimensions[L(j)].width=15
        sh,a,b=info[(tr,t)]
        for k,r in enumerate(range(a,b+1)): ws.cell(2+k,j,f"='{sh}'!D{r}").number_format='0.00'
        last=1+(b-a+1); 
    mx=max(v[2]-v[1]+1 for v in info.values()); base=mx+3
    for lab,off in [('Mean',0),('SEM',1),('n fields',2)]: ws.cell(base+off,1,lab).font=B
    for j in range(1,11):
        col=L(j); n=info[groups[j-1]][2]-info[groups[j-1]][1]+1; rng=f'{col}2:{col}{1+n}'
        ws.cell(base+3,j) # spacer
        ws.cell(base,j,f'=AVERAGE({rng})'); ws.cell(base+1,j,f'=_xlfn.STDEV.S({rng})/SQRT(COUNT({rng}))'); ws.cell(base+2,j,f'=COUNT({rng})')
        ws.cell(base,j).number_format=ws.cell(base+1,j).number_format='0.00'
    ws.cell(base+4,1,'Each value = per-field mean of the cell values on the matching sheet (blank when a field has no value). Cells = nucleus-seeded territories (not segmented cells).').font=Font(name='Calibri',size=10,italic=True)
    for w in wb.worksheets:
        for row in w.iter_rows():
            for x in row:
                if x.value is not None and x.font.name!='Calibri': x.font=N
    wb.save(f'{OUT}/Ecad_{metric}_Timeline.xlsx')
